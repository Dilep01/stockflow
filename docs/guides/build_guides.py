from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
root=Path.cwd(); out=root/'docs/guides'
def base(title,intro):
 d=Document(); sec=d.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.7)
 sec.left_margin=sec.right_margin=Inches(.8)
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  s=d.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.add_paragraph(title,'Title'); d.add_paragraph(intro); d.add_paragraph('Study edition  |  31 October 2026')
 f=sec.footer.paragraphs[0]; f.alignment=2
 field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
 return d
def h(d,t): d.add_heading(t,1)
def p(d,t): d.add_paragraph(t)
def md(d,path):
 text=(root/path).read_text(encoding='utf-8'); fence=False; mermaid=False
 for line in text.splitlines():
  if line.startswith('```'):
   if not fence: mermaid='mermaid' in line
   fence=not fence; continue
  if mermaid and fence: continue
  if not line.strip() or line.startswith('!['): continue
  line=re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',line).replace('**','').replace('`','')
  if line.startswith('#'):
   level=min(len(line)-len(line.lstrip('#')),3); d.add_heading(line.lstrip('# ').replace('—',' '),level)
  elif line.startswith('|'):
   cells=[x.strip() for x in line.strip('|').split('|')]
   if all(re.match(r'^[- :]+$',x) for x in cells): continue
   p(d,'  |  '.join(cells))
  else: p(d,line)
d=base('Stockflow Project Understanding Guide','Learn how the distributed inventory and order management platform works, how its components fit together, and how to explain its implementation. The platform demonstrates local warehouse operations and recovery with synthetic data; it is not a production warehouse deployment.')
h(d,'How to use this guide'); p(d,'Read the business walkthrough first, then architecture and implementation stages. Use the source map to trace the code and the operational procedures to repeat the demo. Test results below are recorded results from 28 and 29 October 2026, not a new test run.')
h(d,'Business walkthrough'); p(d,'A customer chooses products and submits an order. Orders records authoritative prices and asks Inventory to hold stock. Inventory allocates across warehouses, confirms the hold, and Fulfillment creates shipments. Warehouse staff pick and pack each shipment. Dispatch consumes stock before shipment completion is acknowledged. Delivered goods can enter the return inspection flow; sellable goods become available again and damaged goods remain separate.')
p(d,'Example: a warehouse has on_hand 10 and reserved 0. Reserving 3 makes available 7 without changing on_hand. Committing keeps reserved at 3. Dispatch reduces on_hand to 7 and reserved to 0. Returning 1 sellable unit raises on_hand to 8. A damaged return increases damaged stock instead. This example explains the accounting rules; it is not a measured business outcome.')
h(d,'Skills and technologies')
for a,b in [ ('Java 21 and Spring Boot 3.5.16','Backend HTTP APIs, dependency injection, validation, scheduling, transactions and service role selection.'),('Spring JDBC and PostgreSQL','Explicit SQL, row locks, advisory locks, constraints, durable state and ledger reconciliation. JPA is not used.'),('Flyway and Maven','Versioned SQL migrations and repeatable Java dependency resolution and builds.'),('Spring Security and BCrypt','Authentication, password hashing, JWT validation, refresh session rotation and role and ownership checks.'),('Kafka and Spring Kafka','Asynchronous commands and events, consumer groups, partition keys, retries and dead letters.'),('Redis','A 30 second catalog list cache with database fallback; never the inventory authority.'),('React 19.2 and TypeScript 5.9','Component state, forms, role-specific screens and a typed API helper. Broad Row types still use any, so typing could be strengthened.'),('TanStack Query and Vite','Server state fetching and periodic refresh; development server and frontend production bundling.'),('Docker Compose and Nginx','Local service orchestration, routing and authentication rate limiting.'),('Actuator Micrometer Prometheus and Grafana','Health, request and JVM metrics and event backlog monitoring.'),('Mailpit','Captures simulated notification email locally.'),('JUnit Spring Boot tests Vitest and Playwright','Database integration tests, API client tests and browser regression tests. Node HTTP scripts cover concurrency and recovery.'),('OpenAPI JSON Schema Git and PowerShell','API/event contracts, repository tooling and setup, backup, restore and replay scripts. k6 and hosted CI are supplied but recorded execution is not claimed.')]:
 d.add_heading(a,2); p(d,b)
h(d,'Implementation sequence explained')
steps=[('Foundation','Create the Maven backend and Vite frontend, configure PostgreSQL and Compose, and apply common plus role-specific Flyway migrations. Application.java and application.yml define runtime entry and configuration.'),('Identity and access','Implement users, BCrypt passwords, signed access tokens and rotating refresh sessions. Enforce customer ownership and operator warehouse assignment at the backend, not only in the UI.'),('Catalog','Store products and expose catalog operations. Add product snapshots at checkout so later catalog edits cannot rewrite historical order values.'),('Inventory','Define warehouses, balances, reservations and movements. Enforce nonnegative quantities in SQL, lock candidate balances consistently, and record movements with each mutation.'),('Checkout','Validate lines and serialize customer/idempotency-key pairs with an advisory transaction lock. Store order, items, key, history and Reserve event in one transaction.'),('Saga','Process StockReserved, Commit and StockCommitted, then request fulfillment. Persist transitions and outgoing messages so service restarts do not depend on in-memory state.'),('Fulfillment and cancellation','Create warehouse shipments, enforce legal transitions and serialize dispatch against cancellation. Release inventory only after fulfillment acknowledges cancellation.'),('Returns and transfers','Validate return allowance against delivered quantities. Inspect sellable versus damaged goods. Move transfer stock out at dispatch and into the destination only on receipt.'),('Reliable messaging','Add outbox publication, inbox deduplication, bounded retries, dead letters and replay with the original event ID.'),('Distributed runtime','Run six roles from the same image with distinct database owners. Route requests through Nginx; use authenticated catalog lookup and Kafka for domain coordination.'),('Frontend and operations','Build ordering and operations screens, session refresh, query polling, monitoring and backup/restore helpers.'),('Verification','Exercise stock races, duplicate commands, cancellation, returns, outages and browser flows. Record test evidence and retain explicit scope limits.')]
for i,(a,b) in enumerate(steps,1): d.add_heading(f'Step {i} {a}',2); p(d,b)
p(d,'This is the dependency and learning sequence represented by the implementation tracker, not a claimed commit-by-commit chronology.')
for file in ['docs/ARCHITECTURE.md','docs/LEARNING_GUIDE.md','docs/RUNBOOKS.md','docs/VERIFICATION.md','docs/LIMITATIONS.md']:
 d.add_page_break(); md(d,file)
h(d,'Data ownership and source map')
for name,desc in [('Identity','users and refresh_sessions'),('Catalog','products'),('Inventory','warehouses, stock_balances, stock_movements, reservations, reservation_lines, inventory_operations and stock_transfers'),('Orders','orders, order_items, checkout_keys and order_history'),('Fulfillment','fulfillment_orders, shipments, shipment_items, shipment_events and return_requests'),('Notification','notifications'),('Every domain database','outbox_events and processed_messages')]: p(d,name+': '+desc+'.')
p(d,'Domain Java files are under backend/src/main/java/dev/stockflow. SQL definitions are under backend/src/main/resources/db. Orders.java contains checkout and saga transitions; EventWorker.java publishes messages; Dispatcher.java validates and deduplicates delivery in a separate transaction. frontend/src/api.ts manages the in-memory token, shared refresh promise and one retry after HTTP 401. frontend/src/main.tsx defines screens and query polling every five seconds.')
h(d,'Run and demonstrate the project'); p(d,'From the project root in PowerShell 7, run ./scripts/setup.ps1 then ./scripts/start.ps1 -Mode lite. Open http://localhost:8088. For distributed mode use ./scripts/start.ps1 -Mode distributed and http://localhost:8089. Docker Desktop must use Linux containers. Consult README.md for local demo accounts. Never copy demo credentials into a public deployment.')
p(d,'Demonstrate customer checkout, the order timeline, administrator pick/pack/dispatch/delivery, a partial return, a stock transfer, and ledger reconciliation. Use an operator account to demonstrate warehouse restrictions. Lite and distributed databases are separate. Stop with ./scripts/stop.ps1 to retain volumes.')
h(d,'Source references'); p(d,'Repository references: README.md; docs/ARCHITECTURE.md; docs/IMPLEMENTATION.md; docs/LEARNING_GUIDE.md; docs/RUNBOOKS.md; docs/VERIFICATION.md; docs/LIMITATIONS.md; backend/pom.xml; frontend/package.json; backend/src/main/resources/application.yml; backend/src/test/java/dev/stockflow/PlatformTest.java; contracts/*.json. These provide the implementation and recorded verification details used in this guide.')
d.save(out/'Stockflow_Project_Understanding.docx')
qa='''Project and architecture
What is Stockflow?|Stockflow is a local multi-warehouse inventory and order platform. It covers reservation, split fulfillment, cancellation, returns, transfers and recovery using durable database state and events.
What problem does it solve?|Concurrent buyers must not receive the same available stock, and service failures must not silently lose accepted work. The project coordinates these operations and records their state changes.
How would you describe it in an interview?|I would describe the local implementation, then trace checkout through inventory and fulfillment. I would explain one stock race and one recovery case, and identify synthetic tests and production limitations clearly.
Is it a microservices system?|Distributed mode has six separately addressable processes and logically owned databases. They share a build artifact and release lifecycle, so it is not six independently released codebases.
What differs in lite mode?|All domains run in one process and database with transactional local delivery. Distributed mode uses role-specific processes, databases and Kafka. Data is not migrated automatically between modes.
Why use separate databases?|Ownership prevents direct cross-domain database reads and forces explicit contracts. One PostgreSQL server still creates a shared availability dependency.
Why mix HTTP and events?|Checkout needs authoritative product information immediately, so catalog lookup is synchronous. Stock and fulfillment coordination can proceed asynchronously through durable messages.
What does HTTP 202 mean here?|Checkout has accepted and stored work. It does not mean stock was allocated or shipping completed; the client follows the persisted order state.
What are the main limitations?|One broker and database host, shared signing secret, synthetic shipping/refunds, bounded UI lists, and no demonstrated production deployment or multi-host failover.
Java and Spring
Why Java records for request objects?|Records concisely describe request data, while validation annotations constrain fields. They do not make nested collections deeply immutable or replace business validation.
Why use BigDecimal for money?|Decimal arithmetic avoids binary floating point rounding errors in prices. Orders multiplies authoritative product prices by quantities and records the result.
What is dependency injection?|Spring supplies collaborators such as Support, Events and ProductLookup through constructors. This makes dependencies explicit and allows the container to manage their lifecycle.
What does Transactional provide?|Related SQL writes commit or roll back together within the local database. It does not create one transaction across all six services.
Why does Dispatcher use REQUIRES_NEW?|Event consumption has its own transaction, including inbox insertion and business changes. Local publication can redeliver safely if its later bookkeeping fails.
What is a Spring proxy caveat?|Transactional advice normally applies when calls pass through the managed proxy. Self-invocation can bypass that advice, so transaction boundaries must be designed and tested explicitly.
Why JDBC instead of JPA?|Explicit SQL makes locks, conditional updates and transaction scope visible. The tradeoff is manual mapping and more SQL maintenance.
How is input validated?|Bean validation checks request fields; checkout also rejects duplicate SKU lines. Database constraints and domain state checks defend deeper invariants.
Why reject fractional quantities?|The business model counts integer units. Jackson configuration disables accepting floating point values as integers, preventing silent truncation.
How are service roles chosen?|APP_ROLE selects domain components through conditional configuration and selects the corresponding Flyway migration location.
Database and concurrency
What is available stock?|Available equals on_hand minus reserved. Damaged goods are tracked separately and do not increase sellable availability.
Why does a reservation leave on_hand unchanged?|The goods remain physically present. Reserved marks them unavailable to other buyers until dispatch or release.
How is overselling prevented?|PostgreSQL locks serialize competing stock operations, conditional updates and constraints enforce bounds, and reservation and ledger writes share a transaction.
Why not Java synchronized?|It coordinates threads in one JVM only. Two inventory processes need database-level coordination over their shared durable balances.
What is a pessimistic lock?|A row lock prevents conflicting updates while a transaction decides and writes its changes. It favors correctness under contention but can increase waiting.
Why lock in a consistent order?|Acquiring warehouse and SKU rows in a consistent order reduces cycles where transactions each wait for a lock held by the other.
Can locks eliminate every deadlock?|No. Other access paths can still create cycles. Production hardening would monitor deadlocks and use bounded whole-transaction retries where safe.
What is an advisory transaction lock used for?|Checkout locks the customer and idempotency-key combination before testing for an existing request. Concurrent identical submissions cannot independently create orders for that key.
What if advisory hash values collide?|Unrelated requests may wait unnecessarily. The actual customer/key record still determines request identity, so a hash collision is not treated as a duplicate request.
What is the inventory ledger?|An append-only record of stock deltas. A trigger rejects update/delete operations, and corrections are recorded as new movements.
What is reconciliation?|It compares materialized balance fields with ledger totals. An empty result means no discrepancies were found by that comparison.
Why have database constraints if Java validates?|Other code paths and concurrent requests can bypass an earlier assumption. Constraints enforce invariants at the final write boundary.
What does SKIP LOCKED do?|An outbox worker skips rows another worker has locked so workers can claim different messages without waiting on the same row.
What is Flyway used for?|Versioned migrations create and evolve schema predictably. Common messaging migrations combine with either the all-domain set or one service-owned set.
How would you improve query performance?|Measure slow SQL and execution plans first, then examine indexes, rows scanned, lock duration and connection pool waits. Avoid claiming an index improvement without measurements.
Orders and warehouse workflows
What is checkout idempotency?|The same customer and key with the same serialized request returns the existing order. Reusing the key with a changed request is rejected.
Why not trust client prices?|A client can be modified and cached prices can be stale. Checkout reads authoritative product information and snapshots price and product name.
Does equivalent reordered JSON always count as the same request?|The implementation compares the serialized request representation. Do not promise semantic equivalence for differently ordered item lists without explicit normalization.
How is an order allocated?|Inventory first seeks a warehouse that can satisfy all lines. Otherwise it splits by warehouse priority and validates the complete allocation before writing it.
Why validate all allocation lines first?|If one item is unavailable, partial writes must not leave stock held for an order that cannot be fulfilled. Transaction rollback also protects atomicity.
What is a saga?|A persisted sequence of local transactions coordinated by messages. Each service owns its state and uses subsequent actions or compensation to reach a business outcome.
Trace a successful order.|Checkout emits Reserve. StockReserved leads to Commit. StockCommitted confirms the order and triggers CreateFulfillment. FulfillmentReady moves the order into fulfillment.
Do committed reservations expire?|No. Active checkout holds expire, but committed unshipped allocations remain reserved until shipment or cancellation.
How are expiry and commit races handled?|Both use locked reservation state and database time checks. Only a valid active reservation can commit; expired stock cannot be revived by a late confirmation.
Why cancel fulfillment before releasing stock?|If dispatch is already starting, releasing the same allocation could let another buyer consume stock that is leaving the warehouse.
How does dispatch compete with cancellation?|Fulfillment serializes both on its order row. Once dispatch is pending, cancellation is rejected; otherwise acknowledged cancellation permits release.
What is a cancellation tombstone?|A durable record that cancellation happened even if creation has not yet arrived. It blocks delayed creation from resurrecting work.
Why have DISPATCH_PENDING?|Inventory consumption is asynchronous. The intermediate state avoids claiming completed dispatch before inventory acknowledges the stock effect.
How does order status avoid going backward?|Shipment progress uses state ranks and terminal-state guards. Late lower-ranked progress cannot overwrite more advanced order progress.
How do partial returns avoid exceeding shipment quantity?|Allowance checks serialize on the shipment and account for prior non-rejected requests. A rejected request does not consume the remaining return allowance.
How is duplicate restock prevented?|A durable unique operation key identifies the stock effect. Repeating the same operation cannot add the quantity twice.
When does transfer stock become available at destination?|Only on receipt. Dispatch removes source stock and the in-transit quantity is not sellable at either warehouse.
Messaging and recovery
Why use an outbox?|Business changes and event intent commit in the same database transaction. A crash between database commit and broker publication cannot erase that intent.
Why also use an inbox?|The outbox can publish twice. The inbox records the event ID with the business effect so repeated deliveries do not repeat that effect.
Does this provide exactly once delivery?|No. Delivery is at least once, with idempotent local business effects. External SMTP can still duplicate email after a crash.
What if Kafka accepts the event but the publisher crashes?|The outbox may remain pending and publish again. The consumer inbox handles the duplicate event ID.
What if processing fails after inbox insertion?|The transaction rolls back both the inbox entry and business changes. A repaired replay with the original ID can then run.
Why keep the original event ID during replay?|A previously successful event must remain recognizable as a duplicate. Changing its ID can bypass deduplication and repeat an effect.
How are failed publications bounded?|EventWorker increments attempts, delays the next attempt up to 60 seconds, and marks the record dead at 10 attempts. Operators can retry after repairing the cause.
What is a dead letter topic?|A target-specific topic stores messages that exhausted consumer retries. Inspect the event and fix the cause before controlled replay.
What ordering does Kafka guarantee?|Ordering is within a partition, not across the entire system. Aggregate IDs are partition keys; domain guards still handle delayed or conflicting events.
What is a consumer group?|Consumers in the same group share partition work. The two inventory processes use one inventory group and shared inventory database.
Does producer idempotence replace an inbox?|No. It addresses certain broker-side producer duplicates, not every application replay or crash boundary involving database commits.
How are wrong schema versions handled?|Dispatcher rejects unsupported versions and targets not owned by the service before applying business effects.
What happens during a broker outage?|Order intent persists in outbox rows. Delivery resumes when possible; bounded retries can leave dead records requiring an operator retry.
Security
How are passwords stored?|BCrypt hashes passwords. Plaintext passwords should not be stored or logged; documented demo credentials are local-only.
What is in JWT validation?|The application validates signature, HS256 algorithm, issuer, audience and expiry. A valid signature alone is insufficient authorization.
Where are tokens stored?|Access tokens remain in browser memory. Refresh tokens use HttpOnly SameSite Strict cookies, with only hashes stored in the database.
Why rotate refresh tokens?|A successful refresh replaces the prior token under a database lock. Reuse of a replaced token should not create another valid session continuation.
Does logout invalidate every access token immediately?|No. Logout revokes refresh sessions, but existing short-lived access tokens remain usable until expiry without a denylist.
How are object permissions enforced?|Customers are checked against order ownership; operators are checked against warehouse claims. Hiding buttons alone would not protect the API.
What protects refresh against cross-site requests?|The endpoint requires a custom header and rejects cross-site fetches, with strict cookies and no cross-origin CORS policy. These controls do not make XSS harmless.
Why is a shared signing secret a limitation?|Every trusted service holding it can potentially mint tokens. A production design should consider an identity provider or asymmetric signing with controlled key rotation.
Frontend and caching
What does TanStack Query do here?|It manages server data fetching and refetching. The client defaults include one retry, five-second refetching and refetch on window focus.
Why does the UI poll?|Orders complete asynchronously, so the UI needs updated states after acceptance. Polling is simple locally; push updates could be a future improvement.
How does the API helper recover from expiry?|It attempts refresh after a non-auth 401, retries the original request once, and signals session expiry on failure. A shared promise avoids simultaneous refresh calls.
What does TypeScript not guarantee?|Compile-time types do not validate server responses at runtime. The current broad Row type also weakens static guarantees; explicit DTOs would improve maintainability.
Why is cache not the stock authority?|Stale cached stock could permit overselling. Inventory writes use PostgreSQL and checkout reads authoritative product detail.
What happens when Redis fails?|Catalog access falls back to PostgreSQL. The cache is an optimization rather than a mandatory consistency dependency.
Can catalog listings be stale?|Yes, up to the 30-second TTL when invalidation does not take effect. Checkout must recheck authoritative price rather than accept the display value.
Testing operations and tradeoffs
What tests were recorded?|31 backend integration tests, 4 frontend API tests, 3 distributed browser tests, 8 HTTP system scenarios and 4 dead-letter/replay checks. These are historical local results, not production certification.
What did the concurrency experiment show?|100 concurrent requests competed for 10 units; 10 orders reached fulfillment and 90 were rejected. No negative stock or over-reservation was observed in that run.
Can you claim high throughput from that test?|No. Recorded acceptance p95 was 15345 ms and all outcomes took 26682 ms on a loaded local machine. This demonstrates a bounded correctness scenario, not a capacity guarantee.
Why use real PostgreSQL tests?|Row locking, advisory locks, constraints and transaction semantics are database-specific. An in-memory substitute would not establish the same concurrency behavior.
How do tests protect normal demo data?|Backend tests require a separate database ending in _test. System/browser tests create synthetic records in the running demo; recovery scenarios also restart project services.
What do browser tests cover?|Recorded checks include navigation and adjustment, customer checkout/history, invalid login feedback and a 390-pixel mobile viewport without whole-page horizontal overflow.
How would you diagnose a stuck order?|Follow its timeline and correlation ID, inspect domain status and service health, check pending/dead outbox events and Kafka, repair the cause and replay or retry safely.
What should monitoring show?|Request rates, JVM memory, outbox backlog and failures. A backlog that does not drain after recovery deserves investigation; metrics alone do not prove business correctness.
Why is a distributed backup difficult?|Independent databases and Kafka positions must represent a coherent recovery point. The runbook uses a coordinated pause and preserves broker state; isolated logical restore is not full disaster recovery.
What is dangerous about manual stock repair?|Changing balances directly can violate the ledger and obscure the original operation. Reconcile first and use a justified compensating movement.
What would you improve first?|Choose according to deployment goals: replace local security assumptions, add pagination and operational alerts, measure lock contention and load, then plan high availability. These are proposed work, not completed features.
How should you discuss personal contribution?|Describe only the design, implementation, testing or learning work you actually performed. Use the source to explain decisions; do not invent employment experience, customer impact or production ownership.
'''
q=base('Stockflow Interview Questions and Answers','Use these questions to prepare for a project discussion and interviews covering the technologies used in Stockflow. Answers distinguish implemented behavior from proposed improvements. No fixed question bank can cover every interviewer; practice tracing the code and defending the tradeoffs.')
n=0
for line in qa.splitlines():
 if not line: continue
 if '|' not in line: h(q,line); continue
 question,answer=line.split('|',1); n+=1
 q.add_heading(f'{n}. {question}',2); p(q,answer)
h(q,'Practical interview exercises');
for t in ['Trace one order ID across order history, outbox, inbox, reservations and shipments. Explain what should exist after each message.','Sketch two transactions competing for the last unit. Show where the lock is acquired and why only one can reserve it.','Design a duplicate and changed-payload checkout test. Include the same customer/key and two concurrent requests.','Explain a crash immediately after Kafka publication and a crash after SMTP send. Identify why the guarantees differ.','Propose an index using a measured slow query and its execution plan. Explain the write cost and verify the result.','Extend the UI to paginated history. Define the API contract, authorization filter and stable ordering before changing the screen.']:
 p(q,t)
h(q,'Evidence and revision sources'); p(q,'Review docs/ARCHITECTURE.md, docs/VERIFICATION.md, docs/LIMITATIONS.md and docs/RUNBOOKS.md. Trace Orders.java, Inventory.java, Fulfillment.java, EventWorker.java, Dispatcher.java and frontend/src/api.ts. Use PlatformTest.java to locate executable race and rollback scenarios. Recorded tests date to 28 and 29 October 2026; no new application tests were run to create these guides.')
q.save(out/'Stockflow_Interview_Questions_and_Answers.docx'); print('Questions',n)
